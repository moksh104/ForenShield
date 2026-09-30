import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# ── Institute standard constants ─────────────────────────────────────────────
FONT_NAME = 'Times New Roman'
BLACK = RGBColor(0, 0, 0)

def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m_name, m_val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m_name}')
        node.set(qn('w:w'), str(m_val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_table_borders(table, color="000000", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def format_row(row, is_header=False, bg_color=None):
    if is_header:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    for cell in row.cells:
        set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
        if bg_color:
            set_cell_shading(cell, bg_color)
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            for r in p.runs:
                r.font.name = FONT_NAME
                r.font.size = Pt(10 if not is_header else 10.5)
                r.font.color.rgb = BLACK
                if is_header:
                    r.font.bold = True


# ── Heading helpers (Institute style: TNR, black, numbered) ──────────────────

def add_heading_1(doc, text):
    """Chapter-level heading, e.g. 'CHAPTER 1 – INTRODUCTION'"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.name = FONT_NAME
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = BLACK
    return p

def add_heading_2(doc, text):
    """Section heading, e.g. '1.1 Project Summary'"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = FONT_NAME
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = BLACK
    return p

def add_heading_3(doc, text):
    """Sub-section heading, e.g. '1.1.1 Background'"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = FONT_NAME
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = BLACK
    return p


# ── Paragraph / bullet helpers ───────────────────────────────────────────────

def add_para(doc, text, bold_prefix=None, space_after=6,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.5
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = FONT_NAME
        r_b.font.size = Pt(12)
        r_b.font.bold = True
        r_b.font.color.rgb = BLACK
    r = p.add_run(text)
    r.font.name = FONT_NAME
    r.font.size = Pt(12)
    r.font.color.rgb = BLACK
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.5
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = FONT_NAME
        r_b.font.size = Pt(12)
        r_b.font.bold = True
        r_b.font.color.rgb = BLACK
    r = p.add_run(text)
    r.font.name = FONT_NAME
    r.font.size = Pt(12)
    r.font.color.rgb = BLACK
    return p

def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(text)
    r.font.name = FONT_NAME
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.italic = True
    r.font.color.rgb = BLACK
    return p


# ── Image helper ─────────────────────────────────────────────────────────────

def add_image_centered(doc, img_path, width_in_inches=5.2, caption=None):
    if not os.path.exists(img_path):
        print(f"  ⚠ Warning: image not found: {img_path}")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(img_path, width=Inches(width_in_inches))
    if caption:
        add_caption(doc, caption)


# ── Header / Footer (institute format) ───────────────────────────────────────

def add_footer_page_number(footer_para):
    """PPI-CE | <page> | P a g e"""
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_para.paragraph_format.space_after = Pt(0)
    r_prefix = footer_para.add_run("PPI-CE  |  ")
    r_prefix.font.name = FONT_NAME
    r_prefix.font.size = Pt(10)
    r_prefix.font.color.rgb = BLACK

    # Dynamic page-number field
    r_pg = footer_para.add_run()
    r_pg.font.name = FONT_NAME
    r_pg.font.size = Pt(10)
    fld1 = OxmlElement('w:fldChar'); fld1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = "PAGE"
    fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'separate')
    fld3 = OxmlElement('w:fldChar'); fld3.set(qn('w:fldCharType'), 'end')
    r_pg._r.append(fld1); r_pg._r.append(instr); r_pg._r.append(fld2); r_pg._r.append(fld3)

    r_suffix = footer_para.add_run("  |  P a g e")
    r_suffix.font.name = FONT_NAME
    r_suffix.font.size = Pt(10)
    r_suffix.font.color.rgb = BLACK

def add_header_text(header_para,
                    text="ForenShield: Interactive Cybersecurity Training & Digital Investigation Platform"):
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_para.paragraph_format.space_after = Pt(0)
    r = header_para.add_run(text)
    r.font.name = FONT_NAME
    r.font.size = Pt(9)
    r.font.italic = True
    r.font.color.rgb = BLACK
