import os
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from report_utils import (
    FONT_NAME, BLACK,
    add_heading_1, add_heading_2, add_para, add_bullet,
    set_cell_margins, set_cell_shading, set_table_borders, format_row
)


def build_front_matter(doc, img_dir):
    logo_path = os.path.join(img_dir, "be836232e103d9dd5430bb690c9bf9f7d4cf8e1e.jpg")

    # ──────────────────────────────────────────────────────────────────────────
    # COVER PAGE
    # ──────────────────────────────────────────────────────────────────────────
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("PARUL POLYTECHNIC INSTITUTE")
    r.font.name = FONT_NAME; r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = BLACK

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("DEPARTMENT OF COMPUTER ENGINEERING")
    r.font.name = FONT_NAME; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = BLACK

    if os.path.exists(logo_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(16)
        p.add_run().add_picture(logo_path, width=Inches(1.8))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("A PROJECT REPORT ON")
    r.font.name = FONT_NAME; r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = BLACK

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("\u201cFORENSHIELD\u201d")
    r.font.name = FONT_NAME; r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = BLACK

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run("INTERACTIVE CYBERSECURITY TRAINING AND DIGITAL INVESTIGATION PLATFORM")
    r.font.name = FONT_NAME; r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = BLACK

    # Submitted By
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("SUBMITTED BY:")
    r.font.name = FONT_NAME; r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = BLACK

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("MOKSH TARUNKUMAR PATEL (Enrolment No: 2403466160104)")
    r.font.name = FONT_NAME; r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = BLACK

    # Guided By
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("GUIDED BY:")
    r.font.name = FONT_NAME; r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = BLACK

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(20)
    r = p.add_run("PROF. URVASHI PARMAR")
    r.font.name = FONT_NAME; r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = BLACK

    # Degree line
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("DIPLOMA IN COMPUTER ENGINEERING (5th SEMESTER)")
    r.font.name = FONT_NAME; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = BLACK

    # Submitted To
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("SUBMITTED TO:")
    r.font.name = FONT_NAME; r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = BLACK

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("PARUL POLYTECHNIC INSTITUTE\nPOST \u2013 LIMDA, WAGHODIA, VADODARA \u2013 391760\nACADEMIC YEAR: 2025 \u2013 2026")
    r.font.name = FONT_NAME; r.font.size = Pt(11); r.font.color.rgb = BLACK

    doc.add_page_break()

    # ──────────────────────────────────────────────────────────────────────────
    # CERTIFICATE 1 — Guide & HOD (5th Semester)
    # ──────────────────────────────────────────────────────────────────────────
    _add_certificate_header(doc, logo_path)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(36)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(
        "This is to certify that Moksh Tarunkumar Patel, student of Computer Engineering, "
        "Enrolment No: 2403466160104, has satisfactorily completed his project work as a part of "
        "the course curriculum in Diploma in Computer Engineering (5th Semester) having the project title "
        "\u201cFORENSHIELD: INTERACTIVE CYBERSECURITY TRAINING AND DIGITAL INVESTIGATION PLATFORM\u201d "
        "at Parul Polytechnic Institute during the academic year 2025\u20132026."
    )
    r.font.name = FONT_NAME; r.font.size = Pt(12); r.font.color.rgb = BLACK

    # Signature block
    sig_table = doc.add_table(rows=2, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False
    for row in sig_table.rows:
        row.cells[0].width = Inches(3.2)
        row.cells[1].width = Inches(3.2)
        for cell in row.cells:
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

    _sig_cell(sig_table.rows[0].cells[0], "______________________________\nSignature of Guide", WD_ALIGN_PARAGRAPH.LEFT)
    _sig_cell(sig_table.rows[0].cells[1], "______________________________\nHead of Department", WD_ALIGN_PARAGRAPH.RIGHT)
    _sig_cell(sig_table.rows[1].cells[0], "PROF. URVASHI PARMAR\nProject Guide", WD_ALIGN_PARAGRAPH.LEFT, bold=False)
    _sig_cell(sig_table.rows[1].cells[1], "ASST. PROF. POONAM FALDU\nHead of Computer Engineering", WD_ALIGN_PARAGRAPH.RIGHT, bold=False)

    doc.add_page_break()

    # ──────────────────────────────────────────────────────────────────────────
    # CERTIFICATE 2 — Internal & External Jury (5th Semester)
    # ──────────────────────────────────────────────────────────────────────────
    _add_certificate_header(doc, logo_path)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(28)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(
        "This is to certify that Moksh Tarunkumar Patel, student of Computer Engineering, "
        "Enrolment No: 2403466160104, has satisfactorily completed his project work as a part of "
        "the course curriculum in Diploma in Computer Engineering (5th Semester) having the project title "
        "\u201cFORENSHIELD: INTERACTIVE CYBERSECURITY TRAINING AND DIGITAL INVESTIGATION PLATFORM\u201d "
        "and is submitted for project examination."
    )
    r.font.name = FONT_NAME; r.font.size = Pt(12); r.font.color.rgb = BLACK

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(36)
    r = p.add_run("DATE: ____________________\nPLACE: Vadodara")
    r.font.name = FONT_NAME; r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = BLACK

    jury_table = doc.add_table(rows=1, cols=2)
    jury_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    jury_table.autofit = False
    for row in jury_table.rows:
        row.cells[0].width = Inches(3.2)
        row.cells[1].width = Inches(3.2)
        for cell in row.cells:
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

    _sig_cell(jury_table.rows[0].cells[0], "______________________________\nINTERNAL JURY", WD_ALIGN_PARAGRAPH.LEFT)
    _sig_cell(jury_table.rows[0].cells[1], "______________________________\nEXTERNAL JURY", WD_ALIGN_PARAGRAPH.RIGHT)

    doc.add_page_break()

    # ──────────────────────────────────────────────────────────────────────────
    # ACKNOWLEDGEMENT
    # ──────────────────────────────────────────────────────────────────────────
    add_heading_1(doc, "ACKNOWLEDGEMENT")

    add_para(doc,
        "I take this opportunity to express my profound sense of gratitude and respect to all "
        "those who helped me throughout the duration of this project."
    )
    add_para(doc,
        "Firstly, I am extremely grateful to Parul Polytechnic Institute for providing me with an "
        "excellent working environment to undergo my project."
    )
    add_para(doc,
        "I devote my success in this effort to my project guide, Prof. Urvashi Parmar, for "
        "giving me the opportunity to undertake this project and providing crucial feedback "
        "and insightful guidance that influenced me and provided an invaluable opportunity to "
        "undertake and complete the project work in the esteemed institution."
    )
    add_para(doc,
        "I am also deeply thankful to Asst. Prof. Poonam Faldu, Head of the Computer Engineering "
        "Department, whose useful suggestions, gentle soothing attitude, and right directions "
        "helped me a lot to learn in this project, and also for her constant encouragement and "
        "support throughout the project."
    )
    add_para(doc,
        "Last, but not the least, I would like to extend my profound thanks to all my esteemed "
        "colleagues and friends at college level who helped me in the specific areas of this project."
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Yours Sincerely,")
    r.font.name = FONT_NAME; r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = BLACK

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Moksh Tarunkumar Patel\nEnrolment No: 2403466160104")
    r.font.name = FONT_NAME; r.font.size = Pt(12); r.font.color.rgb = BLACK

    doc.add_page_break()

    # ──────────────────────────────────────────────────────────────────────────
    # ABSTRACT
    # ──────────────────────────────────────────────────────────────────────────
    add_heading_1(doc, "ABSTRACT")

    add_para(doc,
        "With the exponential surge in cyber threats, ransomware intrusions, and advanced persistent "
        "threat (APT) campaigns worldwide, there is an urgent demand for skilled cybersecurity practitioners "
        "and digital forensic analysts. Traditional educational methods frequently depend on theoretical "
        "classroom instruction or cumbersome desktop virtualisation setups that lack real-time context and "
        "accessibility for students."
    )
    add_para(doc,
        "To bridge this gap, this project presents ForenShield: Learn. Investigate. Defend.\u2014a cross-platform "
        "mobile cybersecurity training and digital forensics investigation platform. Built using the Flutter "
        "framework with Riverpod state management, and backed by a stateless PHP REST API interfacing with a "
        "cloud-hosted PostgreSQL (Neon) database, ForenShield delivers an immersive, interactive, and gamified "
        "educational environment directly on mobile devices."
    )
    add_para(doc,
        "The system incorporates five core modules: (1) Mission Control, a real-time telemetry dashboard pulling "
        "live vulnerability and threat feeds from CISA Known Exploited Vulnerabilities (KEV), MITRE ATT&CK, and "
        "the National Vulnerability Database (NVD); (2) Cyber Academy, a modular learning curriculum with "
        "interactive quizzes; (3) Simulation Lab, scenario-based guided exercises demonstrating cyber-attack "
        "mechanics in a safe educational environment; (4) Investigation Lab, practical forensic case studies and "
        "OSINT utilities integrated with external intelligence APIs such as VirusTotal; and (5) Gamification "
        "Engine, an XP-based progression system with achievements, streaks, and a global leaderboard."
    )
    add_para(doc,
        "ForenShield establishes an accessible and responsive ecosystem for students and entry-level security "
        "analysts to bridge the divide between theoretical cyber defense concepts and practical digital forensic "
        "investigation skills."
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Keywords: ")
    r.font.name = FONT_NAME; r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = BLACK
    r = p.add_run(
        "Cybersecurity Education, Digital Forensics, Threat Simulation, Flutter, "
        "PHP REST API, PostgreSQL, Gamification, OSINT."
    )
    r.font.name = FONT_NAME; r.font.size = Pt(11); r.font.italic = True; r.font.color.rgb = BLACK

    doc.add_page_break()

    # ──────────────────────────────────────────────────────────────────────────
    # INDEX (Table of Contents)
    # ──────────────────────────────────────────────────────────────────────────
    add_heading_1(doc, "INDEX")

    index_data = [
        ("", "Certificate (5th Semester)", "ii"),
        ("", "Certificate (Jury)", "iii"),
        ("", "Acknowledgement", "iv"),
        ("", "Abstract", "v"),
        ("Chapter 1", "Introduction", "8"),
        ("", "1.1 Project Summary", "8"),
        ("", "1.2 Purpose", "11"),
        ("", "1.3 Scope", "12"),
        ("", "1.4 Benefits", "13"),
        ("Chapter 2", "System Requirement Study", "14"),
        ("", "2.1 Hardware Requirements", "14"),
        ("", "2.2 Software Requirements", "15"),
        ("", "2.3 Functional Requirements", "16"),
        ("", "2.4 Non-Functional Requirements", "18"),
        ("Chapter 3", "System Design / Circuit Diagram", "20"),
        ("", "3.1 System Architecture", "20"),
        ("", "3.2 Flowchart", "21"),
        ("", "3.3 DFD / Use Case / ER Diagram", "23"),
        ("", "3.4 Circuit Diagram (Hardware Abstraction)", "27"),
        ("Chapter 4", "Data Dictionary", "29"),
        ("Chapter 5", "Implementation, Testing & Results", "48"),
        ("", "5.1 Technologies Used", "48"),
        ("", "5.2 Module Description", "49"),
        ("", "5.3 Implementation Details", "50"),
        ("", "5.4 Test Cases", "53"),
        ("", "5.5 Screenshots / Results", "56"),
        ("Chapter 6", "Limitations & Future Enhancement", "77"),
        ("", "6.1 Limitations", "77"),
        ("", "6.2 Future Enhancement", "78"),
        ("Chapter 7", "Conclusion", "80"),
        ("", "References", "82"),
    ]

    idx_table = doc.add_table(rows=len(index_data) + 1, cols=3)
    idx_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    idx_table.autofit = False
    col_widths = [Inches(1.5), Inches(4.2), Inches(0.9)]
    for row in idx_table.rows:
        for c_idx, width in enumerate(col_widths):
            row.cells[c_idx].width = width

    set_table_borders(idx_table, color="000000", sz="4", val="single")

    # Header row
    hdr = idx_table.rows[0]
    hdr.cells[0].paragraphs[0].text = "Sr. No."
    hdr.cells[1].paragraphs[0].text = "Description"
    hdr.cells[2].paragraphs[0].text = "Page No."
    hdr.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_row(hdr, is_header=True, bg_color="D9D9D9")

    for r_idx, (col0, col1, col2) in enumerate(index_data):
        row = idx_table.rows[r_idx + 1]
        p0 = row.cells[0].paragraphs[0]; p0.text = col0
        p1 = row.cells[1].paragraphs[0]; p1.text = col1
        p2 = row.cells[2].paragraphs[0]; p2.text = col2
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        is_chapter = col0.startswith("Chapter") or col0 == "References"
        if is_chapter:
            for c in row.cells:
                for pp in c.paragraphs:
                    for rr in pp.runs:
                        rr.font.bold = True
        format_row(row, is_header=False)

    doc.add_page_break()


# ── Private helpers ──────────────────────────────────────────────────────────

def _add_certificate_header(doc, logo_path):
    """Shared header for both certificate pages."""
    for text, sz in [("PARUL POLYTECHNIC INSTITUTE", 16),
                     ("COMPUTER ENGINEERING DEPARTMENT", 13),
                     ("LIMDA, WAGHODIA, VADODARA \u2013 391760", 11)]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.font.name = FONT_NAME; r.font.size = Pt(sz)
        r.font.bold = (sz > 11); r.font.color.rgb = BLACK

    if os.path.exists(logo_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(14)
        p.add_run().add_picture(logo_path, width=Inches(1.5))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("C E R T I F I C A T E")
    r.font.name = FONT_NAME; r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = BLACK


def _sig_cell(cell, text, align, bold=True):
    p = cell.paragraphs[0]
    p.alignment = align
    r = p.add_run(text)
    r.font.name = FONT_NAME; r.font.size = Pt(11)
    r.font.bold = bold; r.font.color.rgb = BLACK
