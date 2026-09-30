import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from report_utils import (
    FONT_NAME, BLACK,
    add_heading_1, add_heading_2, add_heading_3,
    add_para, add_bullet, add_caption,
    set_cell_margins, set_cell_shading, set_table_borders, format_row
)

def build_chapter_1(doc):
    add_heading_1(doc, "CHAPTER 1 – INTRODUCTION")
    
    add_heading_2(doc, "1.1 Project Summary")
    
    add_heading_3(doc, "1.1.1 Background")
    add_para(
        doc,
        "In today’s interconnected global economy, digital information systems underpin virtually every facet "
        "of human activity, ranging from critical national infrastructure, banking, and healthcare to personal "
        "communications and social media. However, this profound dependence on computational networks has coincided "
        "with an alarming surge in sophisticated cyber threats. Threat actors—ranging from automated botnets and "
        "financially motivated cybercrime syndicates to state-sponsored advanced persistent threats (APTs)—exploit "
        "architectural vulnerabilities, misconfigurations, and human fallibilities on a daily basis. India, "
        "as one of the world's most rapidly expanding digital economies with over 800 million active internet users, "
        "has witnessed exponential increases in phishing, identity theft, unauthorized credential stuffing, "
        "ransomware extortions, and mobile banking scams."
    )
    add_para(
        doc,
        "Despite the escalating frequency and sophistication of these attacks, modern cybersecurity education "
        "in academic institutions continues to be constrained by traditional pedagogical models. A prominent "
        "dilemma confronting computer engineering and information technology departments is the formidable chasm "
        "between theoretical classroom concepts and real-world tactical defense. While students are routinely taught "
        "the mathematical foundations of symmetric cryptography or standard definitions of network protocols, "
        "they rarely gain hands-on operational experience in examining actual attack artifacts, inspecting raw "
        "server logs, tracing forensic evidence across file systems, or responding methodically to live security "
        "breaches. Consequently, university and diploma graduates frequently enter the technology industry ill-equipped "
        "to handle real-world digital investigations and incident triage."
    )
    add_para(
        doc,
        "Therefore, there is an urgent educational requirement for a modern, accessible, and structured mobile "
        "learning platform that bridges this gap. ForenShield has been specifically conceptualized and developed "
        "to fulfill this need. By synthesizing foundational academic learning, realistic threat simulation, and "
        "rigorous digital forensics into an intuitive mobile application, ForenShield delivers a comprehensive "
        "educational environment. The core operational philosophy of the project is grounded in four progressive "
        "phases: Learn, Simulate, Investigate, and Defend."
    )
    
    add_heading_3(doc, "1.1.2 Problem Statement")
    add_para(
        doc,
        "Traditional instructional frameworks for cybersecurity and digital forensics suffer from several systemic "
        "deficiencies that impede effective student learning and long-term skill acquisition:"
    )
    add_bullet(
        doc,
        "Most educational curricula emphasize abstract textbook definitions, mathematical proofs, "
        "and static slides without providing practical exposure to live software vulnerabilities or attack vectors.",
        bold_prefix="Excessive Reliance on Passive Theory: "
    )
    add_bullet(
        doc,
        "Novice students cannot safely practice offensive techniques or observe destructive malware "
        "mechanics on standard computer lab setups due to the inherent risks of compromising institutional local "
        "area networks and computational assets.",
        bold_prefix="Lack of Safe, Sandboxed Simulation Environments: "
    )
    add_bullet(
        doc,
        "Commercial and enterprise forensic suites (such as EnCase, FTK, or Magnet AXIOM) are prohibitively "
        "expensive for academic deployment and feature steep, intimidating learning curves that distract beginners "
        "from grasping fundamental evidence collection methodologies.",
        bold_prefix="Complexity and High Cost of Professional Forensic Suites: "
    )
    add_bullet(
        doc,
        "Conventional homework assignments fail to present digital incidents as unfolding mysteries. "
        "Learners are not trained to correlate disparate pieces of digital evidence (e.g., matching a firewall log entry "
        "with an email header timestamp or a suspicious hash) to reconstruct an objective timeline.",
        bold_prefix="Absence of Evidence Correlation and Investigative Workflows: "
    )
    add_bullet(
        doc,
        "Traditional technical courses often lack motivational incentives, resulting in high drop-off "
        "rates and low completion enthusiasm during self-paced remote study.",
        bold_prefix="Insufficient Student Engagement and Motivation: "
    )
    add_para(
        doc,
        "These deficiencies culminate in a substantial theory-to-practical divide. Students understand what a "
        "SQL injection or phishing attack is in theory, but fail to identify indicators of compromise (IoCs) within "
        "web server logs or formulate an admissible incident report. ForenShield resolves this problem statement "
        "by introducing an integrated, gamified mobile ecosystem that makes advanced cybersecurity education tangible, "
        "interactive, and accessible."
    )
    
    add_heading_3(doc, "1.1.3 Proposed Solution")
    add_para(
        doc,
        "ForenShield offers an innovative, mobile-first unified platform that transforms cybersecurity training from "
        "a passive reading exercise into an active, multi-stage investigative journey. The proposed system provides "
        "a controlled, highly responsive client-server environment where students can seamlessly transition between "
        "theoretical instruction, offensive threat simulation, and defensive digital forensics."
    )
    add_para(
        doc,
        "The end-to-end learning lifecycle supported by ForenShield operates along a structured progression:"
    )
    add_bullet(
        doc,
        "Students access structured, digestible courses categorized by difficulty (Beginner, Intermediate, "
        "Advanced), studying real-world attack concepts followed immediately by formative knowledge-check quizzes.",
        bold_prefix="1. Theory & Foundational Learning (Cyber Academy): "
    )
    add_bullet(
        doc,
        "Learners step into sandboxed, interactive attack scenarios (such as analyzing spoofed phishing emails, "
        "observing malicious SQL injection syntax execution, or witnessing ransomware encryption mechanics) to observe "
        "threat behavior safely from an adversarial perspective.",
        bold_prefix="2. Controlled Threat Simulation (Simulation Lab): "
    )
    add_bullet(
        doc,
        "Students assume the role of junior digital forensic investigators. They are assigned realistic incident "
        "cases, inspect evidence files (packet captures, auth logs, disk artifacts, email headers), correlate timestamps, "
        "and reconstruct incident timelines.",
        bold_prefix="3. Digital Evidence Analysis (Investigation Lab): "
    )
    add_bullet(
        doc,
        "Learners formulate formal forensic verdicts, identifying the root cause, culprit IP/vector, and recommending "
        "mitigation strategies. The system dynamically evaluates their findings against ground truth and generates an "
        "investigation performance score.",
        bold_prefix="4. Verdict Evaluation & Reporting: "
    )
    add_bullet(
        doc,
        "Every successfully solved lesson, simulation scenario, and investigation case awards Experience Points "
        "(XP), unlocks milestone achievement badges, advances student rank tiers (Recruit → Analyst → Specialist → "
        "Investigator → Cyber Sentinel), and updates live institutional leaderboards.",
        bold_prefix="5. Gamified Engagement & Mastery (Mission Control): "
    )
    
    add_heading_3(doc, "1.1.4 Core Functional Pillars")
    add_para(
        doc,
        "The ForenShield architectural design is organized around four core functional pillars:"
    )
    add_bullet(
        doc,
        "A modular educational hub housing multi-chapter courses, rich technical markdown lessons, "
        "and automated multiple-choice quizzes that evaluate theoretical comprehension.",
        bold_prefix="Pillar 1 – Cyber Academy: "
    )
    add_bullet(
        doc,
        "An interactive simulation environment featuring step-by-step attack sandboxes where students "
        "experiment with common cyber fraud scenarios without writing destructive code or breaching network boundaries.",
        bold_prefix="Pillar 2 – Simulation Lab: "
    )
    add_bullet(
        doc,
        "A simulated cybercrime division where students examine case briefs, interrogate digital "
        "evidence items, reconstruct chronological attack sequences, and submit forensic findings.",
        bold_prefix="Pillar 3 – Investigation Lab: "
    )
    add_bullet(
        doc,
        "The operational nerve center providing real-time threat intelligence feeds, active case "
        "trackers, daily mission challenges, XP distribution history, and student performance analytics.",
        bold_prefix="Pillar 4 – Mission Control: "
    )
    
    add_heading_2(doc, "1.2 Purpose")
    add_para(
        doc,
        "The primary purpose of developing ForenShield is to pioneer an accessible, engaging, and pedagogically "
        "sound educational platform that democratizes cybersecurity and digital forensics training for diploma "
        "engineering students. The specific objectives that guided the platform’s engineering include:"
    )
    add_bullet(
        doc,
        "To empower students with a deep, practical understanding of modern cyber threats, vector "
        "mechanisms, and prevention protocols, enabling them to recognize social engineering, malicious URLs, "
        "and credential theft techniques in everyday computing.",
        bold_prefix="Instilling Threat Awareness: "
    )
    add_bullet(
        doc,
        "To replace monotonous textbook memorization with engaging, tactile software interactions where "
        "students learn by doing, experimenting, failing safely, and iterating.",
        bold_prefix="Fostering Experiential Learning: "
    )
    add_bullet(
        doc,
        "To engineer an isolated, completely safe sandbox within the mobile client that demonstrates "
        "how exploits propagate without compromising host operating systems or campus networks.",
        bold_prefix="Safe Threat Simulation: "
    )
    add_bullet(
        doc,
        "To train aspiring engineers in standard digital forensic principles, including the chain of "
        "custody, artifact preservation, log triage, timestamp analysis, and objective hypothesis testing.",
        bold_prefix="Digital Evidence Competence: "
    )
    add_bullet(
        doc,
        "To cultivate critical analytical reasoning by requiring students to synthesize complex, "
        "sometimes contradictory clues, eliminate false positives, and justify their security verdicts.",
        bold_prefix="Analytical Decision-Making: "
    )
    add_bullet(
        doc,
        "To leverage proven behavioral psychology and gamification paradigms (streaks, XP, ranks, "
        "and badges) to sustain learner curiosity and incentivize daily learning habits.",
        bold_prefix="Sustained Student Motivation: "
    )
    
    add_heading_2(doc, "1.3 Scope")
    add_para(
        doc,
        "The functional and operational boundaries of the ForenShield project are well-defined across several "
        "key dimensions:"
    )
    add_bullet(
        doc,
        "Encompasses structured curricula spanning fundamental cybersecurity principles, network "
        "security, web application vulnerabilities (OWASP Top 10), social engineering, digital evidence handling, "
        "and incident response procedures. Automated quizzes validate comprehension at each milestone.",
        bold_prefix="Educational Scope: "
    )
    add_bullet(
        doc,
        "Provides interactive simulation modules for prominent attack vectors: Phishing Email Analysis, "
        "SQL Injection Execution & Prevention, Ransomware Encryption & Recovery Workflows, Man-in-the-Middle (MitM) "
        "Eavesdropping, and Brute-Force Password Cracking. Each simulation provides real-time explanations of "
        "underlying defensive countermeasures.",
        bold_prefix="Simulation Scope: "
    )
    add_bullet(
        doc,
        "Features pre-engineered realistic forensic investigation cases. Students examine multi-modal "
        "artifacts (raw server logs, forensic disk dump summaries, Wireshark PCAP transaction extracts, phishing headers), "
        "construct chronological event timelines, and file official investigative verdicts.",
        bold_prefix="Investigation Scope: "
    )
    add_bullet(
        doc,
        "Implements a mathematical XP progression engine, a dynamic 5-tier rank hierarchy, milestone "
        "achievement badges, streak tracking algorithms, and real-time global leaderboard computations.",
        bold_prefix="Gamification & Analytics Scope: "
    )
    add_bullet(
        doc,
        "Enforces enterprise-grade security protocols, including cryptographic password hashing (bcrypt), "
        "stateless JSON Web Token (JWT) session authorization with short-lived access tokens and refresh tokens, "
        "and email-based One-Time Password (OTP) verification for account recovery.",
        bold_prefix="Authentication & Security Scope: "
    )
    add_bullet(
        doc,
        "Designed as a cross-platform mobile application compiled natively via Flutter for Android "
        "and iOS environments, communicating over secure HTTPS REST interfaces with a cloud-hosted PHP/PostgreSQL "
        "backend and Cloudinary media CDN.",
        bold_prefix="Deployment & Platform Scope: "
    )
    
    add_heading_2(doc, "1.4 Benefits")
    add_para(
        doc,
        "The deployment of ForenShield delivers extensive advantages to students, educators, and academic institutions:"
    )
    add_bullet(
        doc,
        "Transforms abstract cyber concepts into tangible software experiences. Students don't just "
        "read about SQL injection; they see how input manipulation alters database queries and view the resulting logs.",
        bold_prefix="1. Enhanced Conceptual Mastery: "
    )
    add_bullet(
        doc,
        "Improves student vigilance against contemporary phishing schemes, fraudulent links, and "
        "credential compromise, significantly raising personal digital hygiene.",
        bold_prefix="2. Heightened Real-World Cyber Vigilance: "
    )
    add_bullet(
        doc,
        "Eliminates the logistical complexity, legal liability, and network danger of setting up dedicated "
        "physical hacking labs in collegiate environments.",
        bold_prefix="3. Risk-Free Educational Environment: "
    )
    add_bullet(
        doc,
        "Equips diploma students with foundational investigative skills that align directly with entry-level "
        "industry roles such as Junior SOC Analyst, Incident Responder, and Digital Forensic Technician.",
        bold_prefix="4. Career and Industry Readiness: "
    )
    add_bullet(
        doc,
        "Integrates course reading, practical simulations, case investigations, and progress analytics "
        "into a single, unified mobile interface, eliminating fragmented tooling.",
        bold_prefix="5. Centralized and Seamless Learning Experience: "
    )
    add_bullet(
        doc,
        "The underlying backend and client codebases are built upon clean, decoupled service layers, "
        "enabling faculty to effortlessly add new courses, scenarios, and forensic cases without rewriting software logic.",
        bold_prefix="6. Extensibility and Modular Architecture: "
    )
    
    doc.add_page_break()

def build_chapter_2(doc):
    add_heading_1(doc, "CHAPTER 2 – SYSTEM REQUIREMENT STUDY")
    
    add_para(
        doc,
        "A rigorous system requirement study is fundamental to ensuring that software architecture satisfies all "
        "functional, performance, security, and usability benchmarks. This chapter delineates the hardware specifications, "
        "software engineering stack, exhaustive functional requirements, and non-functional quality attributes governing "
        "the ForenShield ecosystem."
    )
    
    add_heading_2(doc, "2.1 Hardware Requirements")
    add_para(
        doc,
        "The development, testing, and operational deployment of ForenShield encompass three hardware tiers: "
        "the developer workstation tier, the target mobile client testing tier, and the backend cloud server infrastructure tier."
    )
    
    # Table 2.1: Developer Workstation Hardware Requirements
    p_t1 = doc.add_paragraph()
    p_t1.paragraph_format.space_before = Pt(8)
    p_t1.paragraph_format.space_after = Pt(4)
    r = p_t1.add_run("Table 2.1: Developer Workstation Hardware Requirements")
    r.font.name = FONT_NAME
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = BLACK
    
    hw_dev_data = [
        ("Component", "Minimum Requirement", "Recommended Specification"),
        ("Processor", "Intel Core i5 (8th Gen) / AMD Ryzen 5", "Intel Core i7 (12th Gen+) / AMD Ryzen 7"),
        ("System RAM", "8 GB DDR4", "16 GB – 32 GB DDR4/DDR5"),
        ("Primary Storage", "256 GB NVMe SSD", "512 GB – 1 TB NVMe M.2 SSD"),
        ("Graphics (GPU)", "Integrated Intel UHD / AMD Vega", "Dedicated NVIDIA GeForce GTX 1650 or higher"),
        ("Display Resolution", "1920 × 1080 Full HD (60 Hz)", "1920 × 1080 or 2560 × 1440 Dual Monitor"),
        ("Network Connectivity", "10 Mbps Broadband Connection", "100 Mbps+ Low-Latency Fiber Internet"),
        ("Peripherals", "Standard USB Keyboard & Mouse", "Ergonomic Keyboard, High-Precision Mouse"),
    ]
    
    t_hw1 = doc.add_table(rows=len(hw_dev_data), cols=3)
    t_hw1.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_hw1.autofit = False
    col_w = [Inches(2.0), Inches(2.2), Inches(2.5)]
    for row in t_hw1.rows:
        for c_idx, w in enumerate(col_w):
            row.cells[c_idx].width = w
    set_table_borders(t_hw1, color="000000", sz="4", val="single")
    
    for r_idx, row_data in enumerate(hw_dev_data):
        row = t_hw1.rows[r_idx]
        is_hdr = (r_idx == 0)
        bg = "D9D9D9" if is_hdr else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            p = row.cells[c_idx].paragraphs[0]
            p.text = val
        format_row(row, is_header=is_hdr, bg_color=bg)
        
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(8)
    p_spacer.paragraph_format.space_after = Pt(4)
    
    # Table 2.2: Mobile Client Hardware Requirements
    p_t2 = doc.add_paragraph()
    p_t2.paragraph_format.space_before = Pt(8)
    p_t2.paragraph_format.space_after = Pt(4)
    r = p_t2.add_run("Table 2.2: Target Mobile Client Device Specifications")
    r.font.name = FONT_NAME
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = BLACK
    
    hw_mob_data = [
        ("Specification", "Minimum Specification", "Recommended Target Specification"),
        ("Operating System", "Android 8.0 (Oreo, API Level 26) / iOS 13", "Android 12.0+ (API Level 31+) / iOS 15+"),
        ("Processor Architecture", "ARMv7 / ARM64 Quad-Core 1.8 GHz", "ARM64 Octa-Core 2.4 GHz+"),
        ("Device Memory (RAM)", "2 GB RAM", "4 GB – 8 GB LPDDR4X RAM"),
        ("Available Storage", "150 MB free internal flash storage", "500 MB+ free internal storage (for cached assets)"),
        ("Display Resolution", "720 × 1280 (HD), 300 ppi", "1080 × 2400 (FHD+), 400+ ppi, 60–120 Hz"),
        ("Touch Interface", "Capacitive multi-touch display", "High-responsiveness capacitive screen"),
        ("Network Hardware", "Wi-Fi 802.11 b/g/n or 4G LTE cellular", "Wi-Fi 5/6 (802.11ac/ax) and 4G/5G LTE"),
    ]
    
    t_hw2 = doc.add_table(rows=len(hw_mob_data), cols=3)
    t_hw2.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_hw2.autofit = False
    for row in t_hw2.rows:
        for c_idx, w in enumerate(col_w):
            row.cells[c_idx].width = w
    set_table_borders(t_hw2, color="000000", sz="4", val="single")
    
    for r_idx, row_data in enumerate(hw_mob_data):
        row = t_hw2.rows[r_idx]
        is_hdr = (r_idx == 0)
        bg = "D9D9D9" if is_hdr else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            p = row.cells[c_idx].paragraphs[0]
            p.text = val
        format_row(row, is_header=is_hdr, bg_color=bg)
        
    add_heading_2(doc, "2.2 Software Requirements")
    add_para(
        doc,
        "ForenShield leverages modern, robust, industry-standard software engineering technologies across its client, "
        "backend, database, and cloud storage tiers:"
    )
    
    # Table 2.3: Software Engineering Technology Stack
    p_t3 = doc.add_paragraph()
    p_t3.paragraph_format.space_before = Pt(8)
    p_t3.paragraph_format.space_after = Pt(4)
    r = p_t3.add_run("Table 2.3: Complete Software Technology Stack")
    r.font.name = FONT_NAME
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = BLACK
    
    sw_data = [
        ("Layer / Component", "Technology / Framework", "Version / Specifications", "Purpose in ForenShield"),
        ("Mobile Client Framework", "Flutter SDK", "3.24.x (Channel Stable)", "Cross-platform high-performance UI rendering"),
        ("Programming Language", "Dart SDK", "3.5.x", "Client application logic, data modeling, type safety"),
        ("State Management", "flutter_riverpod", "2.6.1", "Declarative, compile-safe reactive state management"),
        ("Client Routing", "go_router", "14.6.2", "Declarative URL/path-based navigation & route guards"),
        ("Network HTTP Client", "Dio", "5.10.0", "HTTP/REST client with interceptors for JWT & error retry"),
        ("Local Secure Storage", "flutter_secure_storage", "9.2.2", "Hardware-backed encrypted storage for auth tokens"),
        ("Backend Web Server", "PHP Engine", "8.2.x (with PDO-PGSQL)", "RESTful API processing, JSON serialization, auth"),
        ("Relational Database", "PostgreSQL (Neon Cloud)", "16.x Serverless", "ACID-compliant storage for users, courses, cases"),
        ("Media Asset CDN", "Cloudinary SDK / API", "v2 REST API", "High-performance storage & CDN delivery for avatars"),
        ("Push Notifications", "Firebase Cloud Messaging", "FCM v1 API", "Automated system push alerts and case assignments"),
        ("Integrated Dev Env", "Visual Studio Code / Android Studio", "Latest Stable Releases", "Primary coding, debugging, Android emulation"),
        ("API Testing Tool", "Postman / Thunder Client", "v11.x", "Automated endpoint testing, payload verification"),
        ("Version Control", "Git & GitHub", "2.46.x", "Source code versioning, branch management, CI/CD"),
    ]
    
    t_sw = doc.add_table(rows=len(sw_data), cols=4)
    t_sw.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_sw.autofit = False
    sw_col_w = [Inches(1.6), Inches(1.8), Inches(1.5), Inches(2.1)]
    for row in t_sw.rows:
        for c_idx, w in enumerate(sw_col_w):
            row.cells[c_idx].width = w
    set_table_borders(t_sw, color="000000", sz="4", val="single")
    
    for r_idx, row_data in enumerate(sw_data):
        row = t_sw.rows[r_idx]
        is_hdr = (r_idx == 0)
        bg = "D9D9D9" if is_hdr else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            p = row.cells[c_idx].paragraphs[0]
            p.text = val
        format_row(row, is_header=is_hdr, bg_color=bg)
        
    add_para(
        doc,
        "Detailed descriptions of core software stack selections:",
        bold_prefix="Architectural Rationale: "
    )
    add_bullet(
        doc,
        "Selected over native Android (Java/Kotlin) or React Native due to its high-performance Skia/Impeller "
        "rendering engine, unified single-codebase compiling to native ARM code, and an extensive widget library "
        "that allows pixel-perfect custom cybersecurity UI design.",
        bold_prefix="Flutter & Dart: "
    )
    add_bullet(
        doc,
        "Employed for state management to avoid the boilerplate and runtime exceptions common in "
        "legacy InheritedWidget or Provider approaches. Riverpod offers compile-time safety, seamless asynchronous "
        "data caching, and unidirectional data flow across modules.",
        bold_prefix="Riverpod State Architecture: "
    )
    add_bullet(
        doc,
        "Chosen for the server tier due to its ubiquitous deployment capability, native JSON serialization, "
        "and robust PDO (PHP Data Objects) abstraction that completely mitigates SQL injection risks via parameter "
        "binding when interfacing with the PostgreSQL engine.",
        bold_prefix="PHP 8.2 REST Backend: "
    )
    add_bullet(
        doc,
        "Hosted on the Neon serverless platform, PostgreSQL provides enterprise-grade ACID relational "
        "integrity, complex relational joins across 27 tables, robust foreign key constraint enforcement, and indexed "
        "timestamp querying essential for forensic timeline reconstruction.",
        bold_prefix="PostgreSQL Relational Storage: "
    )
    
    add_heading_2(doc, "2.3 Functional Requirements")
    add_para(
        doc,
        "Functional requirements formally describe the behavioral capabilities, transactions, and features that "
        "the ForenShield system must provide to end users. The requirements are categorized into distinct operational "
        "subsystems as detailed in Table 2.4:"
    )
    
    # Table 2.4: Functional Requirements Matrix
    p_t4 = doc.add_paragraph()
    p_t4.paragraph_format.space_before = Pt(8)
    p_t4.paragraph_format.space_after = Pt(4)
    r = p_t4.add_run("Table 2.4: ForenShield Functional Requirements Matrix")
    r.font.name = FONT_NAME
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = BLACK
    
    fr_data = [
        ("Req ID", "Module", "Functional Requirement Description", "Priority"),
        ("FR-01", "Authentication", "User Registration: New students can register using full name, email, phone, and password with validation.", "High"),
        ("FR-02", "Authentication", "User Login & JWT Issuance: Authenticate credentials using bcrypt; issue short-lived access & refresh tokens.", "High"),
        ("FR-03", "Authentication", "Password Recovery via OTP: Generate 6-digit cryptographic OTP sent to user email with 10-minute expiry.", "High"),
        ("FR-04", "Authentication", "Token Refresh: Silently refresh expired access tokens using valid refresh tokens stored in secure hardware.", "High"),
        ("FR-05", "Cyber Academy", "Course Catalog Listing: Display structured courses categorized by difficulty tier with completion progress.", "High"),
        ("FR-06", "Cyber Academy", "Module & Lesson Delivery: Render lesson content with rich formatting, code blocks, and key takeaways.", "High"),
        ("FR-07", "Cyber Academy", "Quiz Evaluation: Present multiple-choice questions, compute instant scores, record attempts, and award XP.", "High"),
        ("FR-08", "Simulation Lab", "Attack Scenario Catalog: Present categorized threat simulations (Phishing, SQLi, Ransomware, MitM, Brute Force).", "High"),
        ("FR-09", "Simulation Lab", "Interactive Sandbox Execution: Guide user through tactical attack steps, explain mechanics, and demonstrate countermeasures.", "High"),
        ("FR-10", "Simulation Lab", "Simulation Completion Tracking: Record completed scenarios, update user security skill metrics, and award XP.", "Medium"),
        ("FR-11", "Investigation Lab", "Case Briefing & Dossier: Present realistic digital crime cases with background context, victim info, and suspect lists.", "High"),
        ("FR-12", "Investigation Lab", "Evidence Artifact Viewer: Allow inspection of diverse evidence types (server logs, email headers, PCAP, disk dumps).", "High"),
        ("FR-13", "Investigation Lab", "Timeline Reconstruction: Enable chronologically sorting and tagging evidence events to map incident propagation.", "High"),
        ("FR-14", "Investigation Lab", "Verdict Formulation: Allow submitting suspect identification, attack vector, primary evidence, and justification.", "High"),
        ("FR-15", "Investigation Lab", "Automated Verdict Grading: Compare submitted verdict against ground truth, compute accuracy score, and issue case reports.", "High"),
        ("FR-16", "Mission Control", "Command Dashboard: Display active investigation cases, daily streak counters, XP tier, and quick-action shortcuts.", "High"),
        ("FR-17", "Mission Control", "Threat Intelligence Radar: Aggregate and display real-time CVE vulnerability feeds and CISA security advisories.", "Medium"),
        ("FR-18", "Gamification", "XP Engine & Level Progression: Calculate cumulative XP; dynamically transition users across 5 rank designations.", "High"),
        ("FR-19", "Gamification", "Milestone Badges & Achievements: Evaluate unlock triggers upon quiz/case completion and grant collectible badges.", "Medium"),
        ("FR-20", "Gamification", "Global Leaderboard: Compute institutional ranking based on total XP, completed cases, and active streak days.", "Medium"),
        ("FR-21", "Profile & Settings", "User Profile Management: Update profile details, upload avatar to Cloudinary CDN, and change password.", "Medium"),
        ("FR-22", "Security & System", "Session & Audit History: Log user login timestamps, IP addresses, user-agent details, and active device sessions.", "Low"),
    ]
    
    t_fr = doc.add_table(rows=len(fr_data), cols=4)
    t_fr.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_fr.autofit = False
    fr_col_w = [Inches(1.0), Inches(1.8), Inches(3.4), Inches(1.0)]
    for row in t_fr.rows:
        for c_idx, w in enumerate(fr_col_w):
            row.cells[c_idx].width = w
    set_table_borders(t_fr, color="000000", sz="4", val="single")
    
    for r_idx, row_data in enumerate(fr_data):
        row = t_fr.rows[r_idx]
        is_hdr = (r_idx == 0)
        bg = "D9D9D9" if is_hdr else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            p = row.cells[c_idx].paragraphs[0]
            p.text = val
        format_row(row, is_header=is_hdr, bg_color=bg)
        
    add_heading_2(doc, "2.4 Non-Functional Requirements")
    add_para(
        doc,
        "Non-functional requirements specify the operational quality standards, architectural constraints, and performance "
        "thresholds that ensure ForenShield is secure, dependable, responsive, and maintainable:"
    )
    add_bullet(
        doc,
        "All client-server communications must be encrypted in transit using TLS 1.3 over HTTPS. "
        "User passwords must be hashed using the industry-standard bcrypt algorithm (work factor 10) with unique salt. "
        "JWT tokens must be signed with HMAC-SHA256 using a secure 256-bit server secret. Authentication tokens must be "
        "stored locally using Android Keystore / iOS Keychain hardware-backed encryption via flutter_secure_storage. "
        "All database interactions must use parameterized SQL queries through PHP PDO to prevent SQL injection.",
        bold_prefix="1. Security & Data Integrity: "
    )
    add_bullet(
        doc,
        "The mobile application architecture is designed to target smooth 60 frames-per-second (FPS) UI rendering. "
        "REST API design targets response latencies under 350 milliseconds for standard JSON payloads under standard "
        "mobile broadband conditions. Local asset caching ensures fast screen loading.",
        bold_prefix="2. Performance & Response Latency: "
    )
    add_bullet(
        doc,
        "The backend API and cloud PostgreSQL database are architected for high availability during academic "
        "operating hours. The mobile client incorporates robust offline detection and graceful network failure handlers, "
        "preventing application crashes during unexpected connection drops.",
        bold_prefix="3. Reliability & Availability: "
    )
    add_bullet(
        doc,
        "The user interface complies with contemporary Material Design 3 guidelines, featuring clean typography "
        "(Roboto / Inter), high contrast ratios for optimal legibility, consistent iconography, intuitive navigation hierarchies, "
        "and informative visual feedback during all asynchronous operations.",
        bold_prefix="4. Usability & User Experience: "
    )
    add_bullet(
        doc,
        "The codebase adheres to strict modular software engineering paradigms, isolating business logic "
        "into independent feature repositories (Authentication, Academy, Simulation, Investigation, Mission Control). "
        "The serverless PostgreSQL architecture on Neon provides automated horizontal and vertical compute autoscaling to "
        "accommodate concurrent student cohorts during laboratory examination sessions.",
        bold_prefix="5. Maintainability & Scalability: "
    )
    add_bullet(
        doc,
        "The mobile codebase maintains 100% single-source compatibility across both Android (API 26 to 34) "
        "and iOS (iOS 13 to 17) devices without requiring platform-specific forks or modifications.",
        bold_prefix="6. Portability & Cross-Platform Operation: "
    )
    
    doc.add_page_break()
