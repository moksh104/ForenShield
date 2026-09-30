import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION

# Ensure scratch directory is in python module search path
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from report_utils import (
    add_header_text, add_footer_page_number
)
from builder_front_matter import build_front_matter
from builder_chapter1_2 import build_chapter_1, build_chapter_2
from builder_chapter3 import build_chapter_3
from builder_chapter4 import build_chapter_4
from builder_chapter5 import build_chapter_5
from builder_chapter6_7 import build_chapter_6, build_chapter_7, build_references

def create_report():
    print("Initializing ForenShield Document Builder...")
    doc = docx.Document()
    
    # Configure default styles to Times New Roman (Institute standard)
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_after = Pt(6)
    
    # Section 0: Front Matter
    sec0 = doc.sections[0]
    sec0.top_margin = Inches(1.0)
    sec0.bottom_margin = Inches(1.0)
    sec0.left_margin = Inches(1.0)
    sec0.right_margin = Inches(1.0)
    sec0.header.is_linked_to_previous = False
    sec0.footer.is_linked_to_previous = False
    for p in sec0.header.paragraphs:
        p.text = ""
    for p in sec0.footer.paragraphs:
        p.text = ""
        
    img_dir = os.path.abspath(r"C:\Projects\ForenShield\mobile\report_images")
    
    print("Building Front Matter (Cover, Certificates, Acknowledgement, Abstract, Index)...")
    build_front_matter(doc, img_dir)
    
    # Section 1: Main Chapters (Starts with Chapter 1)
    print("Adding Section 2 for Main Chapters with Headers and Footers...")
    sec1 = doc.add_section(WD_SECTION.NEW_PAGE)
    sec1.top_margin = Inches(1.0)
    sec1.bottom_margin = Inches(1.0)
    sec1.left_margin = Inches(1.0)
    sec1.right_margin = Inches(1.0)
    sec1.header.is_linked_to_previous = False
    sec1.footer.is_linked_to_previous = False
    
    # Setup Header & Footer for Chapter Section
    header_p = sec1.header.paragraphs[0]
    add_header_text(header_p, "ForenShield: Interactive Cybersecurity Training & Digital Investigation Platform")
    
    footer_p = sec1.footer.paragraphs[0]
    add_footer_page_number(footer_p)
    
    print("Building Chapter 1 (Introduction)...")
    build_chapter_1(doc)
    
    print("Building Chapter 2 (System Requirement Study)...")
    build_chapter_2(doc)
    
    print("Building Chapter 3 (System Design / Circuit Diagram)...")
    build_chapter_3(doc, img_dir)
    
    print("Building Chapter 4 (Data Dictionary)...")
    build_chapter_4(doc)
    
    print("Building Chapter 5 (Implementation, Testing & Results)...")
    build_chapter_5(doc, img_dir)
    
    print("Building Chapter 6 (Limitations & Future Enhancement)...")
    build_chapter_6(doc)
    
    print("Building Chapter 7 (Conclusion)...")
    build_chapter_7(doc)
    
    print("Building References...")
    build_references(doc)
    
    output_path = os.path.abspath(r"C:\Projects\ForenShield\mobile\ForenShield_Final_Project_Report.docx")
    doc.save(output_path)
    print(f"Report generated successfully and saved to: {output_path}")

if __name__ == "__main__":
    create_report()
