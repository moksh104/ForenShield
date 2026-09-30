import docx

doc = docx.Document(r'C:\Projects\ForenShield\mobile\ForenShield_Final_Project_Report.docx')

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")
print(f"Total Sections: {len(doc.sections)}")

word_count = 0
for p in doc.paragraphs:
    word_count += len(p.text.split())
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            word_count += len(c.text.split())

print(f"Total Words (approx): {word_count}")

# Check fonts used across document
fonts_used = set()
font_sizes = set()
for p in doc.paragraphs:
    for r in p.runs:
        if r.font.name:
            fonts_used.add(r.font.name)
        if r.font.size:
            font_sizes.add(r.font.size.pt)
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    if r.font.name:
                        fonts_used.add(r.font.name)
                    if r.font.size:
                        font_sizes.add(r.font.size.pt)

print(f"Fonts detected: {sorted(list(fonts_used))}")
print(f"Font sizes detected: {sorted(list(font_sizes))}")

# Check images embedded
drawings_count = 0
for p in doc.paragraphs:
    drawings = p._element.xpath('.//w:drawing')
    drawings_count += len(drawings)
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            drawings = c._element.xpath('.//w:drawing')
            drawings_count += len(drawings)

print(f"Total embedded drawings/images: {drawings_count}")

# Check Certificate text
cert_text = []
for p in doc.paragraphs[:35]:
    if "certify" in p.text.lower():
        cert_text.append(p.text)

print(f"Certificate texts found: {len(cert_text)}")
for ct in cert_text:
    print(f"  --> {ct[:120]}...")
    if "5th semester" in ct.lower():
        print("      [CHECK PASSED] Explicitly mentions 5th Semester!")
    if "6th semester" in ct.lower():
        print("      [WARNING] Mentions 6th semester!")

# Check meta-language
banned_phrases = [
    "as per the guidelines",
    "the institute calls this chapter",
    "in this report we will",
    "note to the reviewer",
    "ai generated",
    "placeholder",
    "lorem ipsum"
]
found_banned = []
for i, p in enumerate(doc.paragraphs):
    for bp in banned_phrases:
        if bp in p.text.lower():
            found_banned.append((i, bp, p.text[:60]))

print(f"Banned meta-language matches: {len(found_banned)}")
for fb in found_banned:
    print(f"  P{fb[0]}: matched '{fb[1]}' -> {fb[2]}")
