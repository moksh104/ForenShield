import docx

doc = docx.Document(r'C:\Projects\ForenShield\mobile\ForenShield_Final_Project_Report.docx')

with open(r'C:\Projects\ForenShield\mobile\scratch_images_map.txt', 'w', encoding='utf-8') as out:
    for i, p in enumerate(doc.paragraphs):
        drawings = p._element.xpath('.//w:drawing')
        if drawings:
            next_p = doc.paragraphs[i+1].text.strip() if i+1 < len(doc.paragraphs) else ''
            prev_p = doc.paragraphs[i-1].text.strip() if i > 0 else ''
            blips = p._element.xpath('.//a:blip/@r:embed')
            out.write(f"P{i:03d}: drawings={len(drawings)}, blips={blips}\n")
            out.write(f"   text: {p.text.strip()[:60]}\n")
            out.write(f"   prev: {prev_p[:60]}\n")
            out.write(f"   next: {next_p[:60]}\n\n")

print("Done writing scratch_images_map.txt")
